"""Integrated Excel workbook of one request (D-8, D-12.9–D-12.13).

Built on demand and never stored (D-3.3). First sheet: the status of every document; then the
six v1 sheets with the rows of all documents, each row prefixed with request/document id,
company and file name. Values that Excel can't hold are adjusted in the workbook only.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from core.db import session_scope
from core.fa import fa_digits, jalali_datetime
from core.models import SHEET_TABLES, AnalysisResult, Document, UserDocument, as_utc
from core.sheets import SHEETS
from services import results
from services.jobs import FAILED, REUSED, SUCCEEDED, DocState, RequestState, doc_explanation, doc_status_label

STATUS_SHEET = "وضعیت_اسناد"
ID_COLUMNS = ("شناسه_درخواست", "شناسه_سند", "نام_شرکت", "نام_فایل")
EXTRA_COLUMN = "سایر_فیلدها"
STATUS_COLUMNS = (
    "ردیف", "شناسه_درخواست", "شناسه_سند", "نام_فایل", "نام_شرکت", "سال", "وضعیت", "توضیح", "ارائه‌دهنده", "مدل",
    "منبع_نتیجه", "زمان_تحلیل", "زمان_شروع", "زمان_پایان", "تعداد_تلاش", "مدت_آخرین_تلاش_ثانیه",
    "مجموع_زمان_پردازش_ثانیه", "توکن_ورودی", "توکن_خروجی", "مجموع_توکن_ورودی", "مجموع_توکن_خروجی",
    "دما_اعمال_شد", "ذخیره_در_پایگاه", "PDF_نگه‌داری_شد", "SHA-256",
)
NO_TOKENS_NOTE = "تلاش‌هایی که پاسخی دریافت نکردند (مهلت یا خطای سرور) توکنی گزارش نمی‌کنند."
CELL_LIMIT = 32_767
CUT_MARKER = "… متن کامل در JSON"
# v1's long-text columns get a capped width (style_excel_file)
LONG_TEXT = ("نکات_کلیدی_و_نتیجه_گیری", "جزئیات", "شرح", "خلاصه_دلایل", "جزییات_سطح_ریسک_تعیین_شده", "توضیح",
             EXTRA_COLUMN)
FONT = "Vazirmatn"
HEADER_FILL = PatternFill(start_color="CCC0DA", end_color="CCC0DA", fill_type="solid")
ALT_FILL = PatternFill(start_color="F7F7F7", end_color="F7F7F7", fill_type="solid")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# Status colours (Excel's good / neutral / bad styles), number formats and print layout (backlog Q2)
TONES = {
    "success": (PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"), "006100"),
    "warning": (PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"), "9C5700"),
    "danger": (PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"), "9C0006"),
}
NUMBER_FORMATS = {
    "توکن_ورودی": "#,##0", "توکن_خروجی": "#,##0", "مجموع_توکن_ورودی": "#,##0", "مجموع_توکن_خروجی": "#,##0",
    "مدت_آخرین_تلاش_ثانیه": "0.0", "مجموع_زمان_پردازش_ثانیه": "0.0",
}


def workbook_name(state: RequestState) -> str:
    return f"{state.id}.xlsx"


@dataclass
class _Doc:
    doc: DocState
    rows: dict[str, list[dict]] | None  # sheet key -> rows (db column names); None = no result available
    provider: str
    model: str
    analysed_at: object
    notes: list[str] = field(default_factory=list)  # cells adjusted in Excel (D-12.12)


def _collect(state: RequestState) -> list[_Doc]:
    """Result rows and metadata of every document, from the DB or from memory."""
    collected = []
    with session_scope() as db:
        for doc in state.docs:
            provider, model, analysed_at = state.provider_name, state.model_name, doc.finished_at
            rows = None
            if doc.status in (SUCCEEDED, REUSED) and doc.stored and doc.result_request_id:
                result = db.get(AnalysisResult, (doc.result_request_id, doc.document_id))
                if result is not None:
                    provider, model, analysed_at = result.provider_name, result.model_name, as_utc(result.created_at)
                    rows = {}
                    for spec in SHEETS:
                        table = SHEET_TABLES[spec.key]
                        rows[spec.key] = [dict(r._mapping) for r in db.execute(
                            select(table).where(table.c.request_id == doc.result_request_id,
                                                table.c.document_id == doc.document_id).order_by(table.c.row_no))]
            elif doc.status == SUCCEEDED and doc.data is not None:
                rows = results.prepare(doc.data).rows
            collected.append(_Doc(doc, rows, provider, model, analysed_at))
    return collected


def _kept_pdfs(state: RequestState) -> set[str]:
    if not state.persist:
        return set()
    with session_scope() as db:
        return set(db.scalars(select(UserDocument.document_id).join(Document, Document.id == UserDocument.document_id)
                              .where(UserDocument.user_id == state.user_id, Document.storage_path.is_not(None))))


def _cell(value, note: list[str] | None = None, where: str = ""):
    """Excel-safe value: no control characters, at most 32,767 characters (D-12.12)."""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(value) > CELL_LIMIT:
        value = value[:CELL_LIMIT - len(CUT_MARKER) - 1] + " " + CUT_MARKER
        if note is not None:
            note.append(where)
    return value


def _yes_no(value: bool | None) -> str:
    return "" if value is None else ("بله" if value else "خیر")


def _is_long(header: str) -> bool:
    return any(name in str(header) for name in LONG_TEXT)


def _style(ws, headers: tuple[str, ...] | list[str]) -> None:
    """v1 look (header colour, borders, alternating rows, widths), right-to-left, frozen header, filters.

    Long text reads from the right edge and the top of the cell; numbers get formats; printing is
    landscape, one page wide, with the header row on every page.
    """
    ws.sheet_view.rightToLeft = True
    base_font, header_font = Font(name=FONT, size=11), Font(name=FONT, size=11, bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    paragraph = Alignment(horizontal="right", vertical="top", wrap_text=True)
    long = [_is_long(h) for h in headers]
    formats = [NUMBER_FORMATS.get(str(h)) for h in headers]
    widths = [len(str(h)) for h in headers]
    for row_index, row in enumerate(ws.iter_rows(), start=1):
        for col_index, cell in enumerate(row):
            cell.border = BORDER
            if row_index == 1:
                cell.font, cell.fill, cell.alignment = header_font, HEADER_FILL, center
            else:
                cell.font = base_font
                cell.alignment = paragraph if long[col_index] else center
                if formats[col_index]:
                    cell.number_format = formats[col_index]
                if row_index % 2 == 0:
                    cell.fill = ALT_FILL
                if cell.value is not None:
                    widths[col_index] = max(widths[col_index], min(len(str(cell.value)), 200))
    for index in range(len(headers)):
        width = min(max(widths[index], 20), 40) if long[index] else min(max(widths[index] + 3, 15), 50)
        ws.column_dimensions[get_column_letter(index + 1)].width = width
    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def status_tone(doc: DocState) -> str:
    """green = result available, red = failed, orange = anything else (warnings, skipped, stopped, duplicate)."""
    if doc.status in (SUCCEEDED, REUSED) and not doc.warnings:
        return "success"
    return "danger" if doc.status == FAILED else "warning"


def _paint_status(ws, tones: list[str]) -> None:
    column = STATUS_COLUMNS.index("وضعیت") + 1
    for row_index, tone in enumerate(tones, start=2):
        fill, color = TONES[tone]
        cell = ws.cell(row_index, column)
        cell.fill, cell.font = fill, Font(name=FONT, size=11, bold=True, color=color)


def _sheet_value(column: str, value):
    if column == "is_applicable" and value in ("True", "False"):
        return value == "True"  # Python-produced flag; v1 wrote a boolean
    return value


def build_workbook(state: RequestState) -> bytes:
    docs = _collect(state)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for spec in SHEETS:
        with_extra = any(row.get("extra") for d in docs if d.rows for row in d.rows.get(spec.key, []))
        headers = [*ID_COLUMNS, *(persian for _, persian in spec.columns), *([EXTRA_COLUMN] if with_extra else [])]
        ws = workbook.create_sheet(spec.title)
        ws.append(headers)
        for d in docs:
            if not d.rows:
                continue
            summary = d.doc.summary or {}
            for row in d.rows.get(spec.key, []):
                line = [state.id, d.doc.document_id[:12], _cell(summary.get("company_name")), _cell(d.doc.file_name)]
                for column, persian in spec.columns:
                    where = f"{spec.title} › {persian} (ردیف {fa_digits(ws.max_row)})"
                    line.append(_cell(_sheet_value(column, row.get(column)), d.notes, where))
                if with_extra:
                    extra = row.get("extra")
                    text = json.dumps(extra, ensure_ascii=False) if extra else None
                    line.append(_cell(text, d.notes, f"{spec.title} › {EXTRA_COLUMN} (ردیف {fa_digits(ws.max_row)})"))
                ws.append(line)
        _style(ws, headers)

    kept = _kept_pdfs(state)
    ws = workbook.create_sheet(STATUS_SHEET, 0)
    ws.append(STATUS_COLUMNS)
    tones = []
    for number, d in enumerate(docs, start=1):
        doc, summary = d.doc, d.doc.summary or {}
        explanation = doc_explanation(doc)
        if doc.status in (SUCCEEDED, REUSED) and d.rows is None:
            explanation = " ".join(filter(None, [explanation, "نتیجه این سند دیگر در دسترس نیست."]))
        if doc.status == FAILED and doc.attempts and not doc.total_input_tokens:
            explanation = " ".join(filter(None, [explanation, NO_TOKENS_NOTE]))
        if d.notes:
            explanation = " ".join(filter(None, [explanation, "متن این خانه‌ها در اکسل کوتاه شد (متن کامل در JSON): "
                                                 + "، ".join(d.notes)]))
        has_result = doc.status in (SUCCEEDED, REUSED)
        ws.append([
            number, state.id, doc.document_id[:12], _cell(doc.file_name), _cell(summary.get("company_name")),
            summary.get("year"), doc_status_label(doc), _cell(explanation),
            d.provider if has_result else state.provider_name, d.model if has_result else state.model_name,
            doc.result_request_id or "", jalali_datetime(d.analysed_at) if has_result and d.analysed_at else "",
            jalali_datetime(doc.started_at) if doc.started_at else "",
            jalali_datetime(doc.finished_at) if doc.finished_at else "",
            doc.attempts or None, doc.last_call_seconds, doc.total_call_seconds or None, doc.input_tokens,
            doc.output_tokens, doc.total_input_tokens or None, doc.total_output_tokens or None,
            _yes_no(doc.temperature_applied),
            _yes_no(doc.stored) if has_result else "", _yes_no(doc.document_id in kept) if state.persist else "خیر",
            doc.document_id,
        ])
        tones.append("warning" if has_result and d.rows is None else status_tone(doc))
    _style(ws, STATUS_COLUMNS)
    _paint_status(ws, tones)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
