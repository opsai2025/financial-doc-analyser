"""Word report of one request (D-16.1, D-17.4–D-17.8).

Built on demand from the same rows as the Excel workbook and never stored (D-3.3). A4 portrait,
right-to-left, B Nazanin: a cover with the request's details and a summary table, then one section
per company (grouped as on the dashboard) in which every topic lists the items of all its
documents together, in chronological order, each row labelled with its document («سند»).
Items of sections marked "not applicable" and checklist rows not found in the report are left
out (they stay in Excel); no charts.
"""

from __future__ import annotations

import io
import unicodedata
from collections import Counter
from dataclasses import dataclass

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from core.fa import fa_digits, jalali_datetime
from core.models import utcnow
from services.dashboard import applicable, company_key
from services.excel import DocRows, collect_docs
from services.jobs import RequestState, doc_explanation, doc_status_label

FONT = "B Nazanin"  # D-17.4; Word substitutes where it is not installed
SIZES = {"body": 12.5, "cell": 11.5, "h1": 19, "h2": 16, "h3": 13.5, "footer": 10}  # D-16.1 sizes + 2 pt
NAVY = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x52, 0x60, 0x7A)
HEADER_FILL = "CCC0DA"  # the Excel header colour (D-16.2 keeps the v1 look)
TITLE = "گزارش تحلیل گزارش‌های حسابرسی"
NOT_IN_EXCEL_NOTE = ("متن‌ها همان پاسخ مدل است. مواردی که مدل «موضوعیت ندارد» تشخیص داده و موضوعاتی از چک‌لیست که "
                     "در گزارش حسابرس نیامده‌اند، در این گزارش نیستند و در فایل اکسل همین درخواست موجودند. "
                     "ستون «سند» هر ردیف را به سند آن (جدول «اسناد» هر شرکت) وصل می‌کند.")
NONE_FOUND = "موردی ثبت نشده است."
NO_YEAR = "بدون سال"
LRE, RLE, POP = "\u202a", "\u202b", "\u202c"  # left/right-to-left embedding, pop
DOC_COLUMN = ("سند", 1.9)  # first column of every company table (D-17.6)

# Item tables per sheet after the «سند» column: (heading, [(label, columns joined with " / ", width in cm)])
ITEM_TABLES = (
    # basis items have no paragraph/page reference in the schema, so that (always empty) column is left out
    ("basis", "مبانی اظهارنظر", (("شماره", ("item_no",), 1.5), ("عنوان", ("title",), 3.4), ("شرح", ("description",), 7.6),
                                 ("نوع دلیل", ("reason_type",), 2.7))),
    ("emphasis", "تأکید بر مطالب خاص", (("عنوان", ("title",), 2.9), ("شرح", ("description",), 5.0),
                                        ("ریسک برجسته‌شده", ("highlighted_risk",), 2.6),
                                        ("دسته / زیرشاخه", ("risk_category", "risk_subcategory"), 2.6),
                                        ("بند / صفحه", ("paragraph_no", "page_no"), 2.0))),
    ("legal", "گزارش قانونی (تخلفات)", (("عنوان تخلف", ("violation_title",), 2.9), ("شرح", ("description",), 4.6),
                                        ("مبانی قانونی", ("legal_bases",), 3.0),
                                        ("دسته / زیرشاخه", ("category", "subcategory"), 2.6),
                                        ("بند / صفحه", ("paragraph_no", "page_no"), 2.0))),
    ("checklist", "چک‌لیست (موضوعات آمده در گزارش)", (("موضوع", ("topic",), 3.4), ("وضعیت", ("status",), 2.8),
                                                      ("جزئیات", ("details",), 6.9),
                                                      ("بند / صفحه", ("paragraph_no", "page_no"), 2.0))),
)


def keep_row(sheet: str, row: dict) -> bool:
    """Items shown in the report: not the placeholder of a section marked not applicable (D-17.3);
    checklist topics only when «در گزارش آمده» is true (D-17.7)."""
    if sheet == "checklist":
        return _text(row.get("in_report")) == "True"
    return applicable(row)


def report_name(state: RequestState) -> str:
    return f"{state.id}.docx"


def _text(value) -> str:
    """Word-safe text: no control characters (python-docx rejects them)."""
    if value is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", value if isinstance(value, str) else str(value)).strip()


def _file(name) -> str:
    """A file name in its own direction, so "1400.pdf" isn't shown as "pdf.1400" (as D-12.16 in the app).

    Word doesn't reliably support the isolate marks the app uses; an embedding (LRE/RLE … PDF) chosen
    from the first strong character works in every version.
    """
    name = _text(name)
    for char in name:
        kind = unicodedata.bidirectional(char)
        if kind in ("R", "AL"):
            return f"{RLE}{name}{POP}"
        if kind == "L":
            break
    return f"{LRE}{name}{POP}" if name else ""


# --- Right-to-left formatting (python-docx has no API for most of it) -------------------------

def _insert(parent, tag: str, successors: tuple[str, ...]):
    """The child `tag` of `parent`, created in schema order when missing."""
    found = parent.find(qn(tag))
    if found is None:
        found = OxmlElement(tag)
        parent.insert_element_before(found, *successors)
    return found


RPR_AFTER_SZCS = ("w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
                  "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
PPR_AFTER_BIDI = ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _font(rpr, *, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    """The report font for every script, right-to-left; size and bold also for complex scripts (Persian)."""
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.get_or_add_rtl().set(qn("w:val"), "1")
    if bold is not None:
        rpr.get_or_add_b().set(qn("w:val"), "1" if bold else "0")
        rpr.get_or_add_bCs().set(qn("w:val"), "1" if bold else "0")
    if color is not None:
        rpr.get_or_add_color().set(qn("w:val"), str(color))
    if size is not None:
        half_points = str(int(size * 2))
        rpr.get_or_add_sz().set(qn("w:val"), half_points)
        _insert(rpr, "w:szCs", RPR_AFTER_SZCS).set(qn("w:val"), half_points)


def _bidi(ppr) -> None:
    _insert(ppr, "w:bidi", PPR_AFTER_BIDI)


def _style(document, name: str, *, size: float, bold: bool = False, color: RGBColor | None = None,
           before: float = 0, after: float = 4):
    style = document.styles[name]
    _font(style.element.get_or_add_rPr(), size=size, bold=bold, color=color)
    _bidi(style.element.get_or_add_pPr())
    style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
    return style


def _new_document():
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.0)
    _style(document, "Normal", size=SIZES["body"], after=4)
    _style(document, "Heading 1", size=SIZES["h1"], bold=True, color=NAVY, after=8)
    _style(document, "Heading 2", size=SIZES["h2"], bold=True, color=NAVY, before=12, after=6)
    _style(document, "Heading 3", size=SIZES["h3"], bold=True, color=NAVY, before=10, after=4)
    return document


def _paragraph(document, text: str, *, style: str | None = None, muted: bool = False, bold: bool = False):
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    if muted or bold:
        _font(run._r.get_or_add_rPr(), color=MUTED if muted else None, bold=bold or None)
    return paragraph


def _cell(cell, text: str, *, header: bool = False, size: float = SIZES["cell"]) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    _font(run._r.get_or_add_rPr(), size=size, bold=header or None)
    if header:
        tcpr = cell._tc.get_or_add_tcPr()
        shading = _insert(tcpr, "w:shd", ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                                          "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
                                          "w:tcPrChange"))
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), HEADER_FILL)


def _table(document, headers: list[str] | None, rows: list[list[str]], widths_cm: list[float]):
    """A bordered right-to-left table (first column on the right); the header row repeats on every page."""
    table = document.add_table(rows=0, cols=len(widths_cm))
    table.style = document.styles["Table Grid"]
    table.autofit = False
    _insert(table._tbl.tblPr, "w:bidiVisual", ("w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
                                                "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                                                "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                                "w:tblDescription", "w:tblPrChange"))
    lines = ([headers] if headers else []) + rows
    for index, line in enumerate(lines):
        row = table.add_row()
        header = headers is not None and index == 0
        if header:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, line, widths_cm, strict=True):
            cell.width = Cm(width)
            _cell(cell, value, header=header)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _key_values(document, pairs: list[tuple[str, str]]) -> None:
    table = _table(document, None, [[label, value] for label, value in pairs if value], [4.0, 13.0])
    for row in table.rows:
        _font(row.cells[0].paragraphs[0].runs[0]._r.get_or_add_rPr(), bold=True)


# --- Content ---------------------------------------------------------------------------------

def _summary(d: DocRows) -> dict:
    return ((d.rows or {}).get("summary") or [{}])[0]


def _company_name(d: DocRows) -> str:
    return _text(_summary(d).get("company_name") or (d.doc.summary or {}).get("company_name"))


def _year(d: DocRows) -> int | None:
    year = _summary(d).get("year") or (d.doc.summary or {}).get("year")
    try:
        return int(year) if year is not None else None
    except (TypeError, ValueError):
        return None


@dataclass
class _Company:
    name: str
    docs: list[DocRows]  # chronological (D-17.8)
    labels: dict[int, str]  # id(DocRows) -> «سند» label


def _companies(docs: list[DocRows], merge_map: dict[str, str]) -> list[_Company]:
    """Documents grouped like the dashboard: Level-B company key plus the user's merges (D-17.5)."""
    groups: dict[str, list[tuple[str, DocRows]]] = {}
    for d in docs:
        own_key = company_key(_company_name(d)) or f"?{d.doc.document_id[:12]}"
        groups.setdefault(merge_map.get(own_key, own_key), []).append((own_key, d))
    companies = []
    for group, members in groups.items():
        own = [_company_name(d) for key, d in members if key == group and _company_name(d)]
        name = Counter(own or [_company_name(d) for _, d in members if _company_name(d)] or ["نامشخص"]).most_common(1)[0][0]
        ordered = sorted((d for _, d in members),
                         key=lambda d: (_year(d) is None, _year(d) or 0, _text(d.doc.file_name), d.doc.position))
        companies.append(_Company(name, ordered, _labels(ordered)))
    return sorted(companies, key=lambda c: c.name)


def _labels(ordered: list[DocRows]) -> dict[int, str]:
    """«سند» label: the year, numbered when a company has several documents of that year (D-17.6)."""
    per_year = Counter(_year(d) for d in ordered)
    seen: Counter = Counter()
    labels = {}
    for d in ordered:
        year = _year(d)
        base = fa_digits(year) if year is not None else NO_YEAR
        seen[year] += 1
        labels[id(d)] = f"{base} ({fa_digits(seen[year])})" if per_year[year] > 1 else base
    return labels


def _joined(row: dict, columns: tuple[str, ...]) -> str:
    return " / ".join(t for t in (_text(row.get(c)) for c in columns) if t)


def _cover(document, state: RequestState, companies: list[_Company], without: list[DocRows]) -> None:
    title = document.add_paragraph(TITLE, style="Heading 1")
    title.paragraph_format.alignment = 1  # center
    with_result = [d for c in companies for d in c.docs]
    counts = [f"{fa_digits(len(with_result))} سند با نتیجه از {fa_digits(len(companies))} شرکت"]
    if without:
        counts.append(f"{fa_digits(len(without))} سند بدون نتیجه")
    models = sorted({f"{d.provider} › {d.model}" for d in with_result}) or [f"{state.provider_name} › {state.model_name}"]
    _key_values(document, [
        ("شناسه درخواست", state.id),
        ("کاربر", _text(state.username)),
        ("زمان ایجاد درخواست", jalali_datetime(state.created_at)),
        ("ارائه‌دهنده › مدل", "، ".join(_text(m) for m in models)),
        ("اسناد", "، ".join(counts)),
        ("زمان تهیه گزارش", jalali_datetime(utcnow())),
    ])

    document.add_paragraph("خلاصه اسناد", style="Heading 2")
    if with_result:
        rows, number = [], 0
        for company in companies:
            for d in company.docs:
                number += 1
                rows.append([fa_digits(number), company.name, company.labels[id(d)],
                             _text(_summary(d).get("opinion_type")), _text(_summary(d).get("overall_risk")),
                             _file(d.doc.file_name)])
        _table(document, ["ردیف", "شرکت", "سند", "نوع اظهارنظر", "سطح ریسک کلی", "فایل"], rows,
               [1.2, 4.4, 1.9, 2.8, 2.5, 4.2])
    else:
        _paragraph(document, "هیچ سندی از این درخواست نتیجه ندارد.", muted=True)

    if without:
        document.add_paragraph("اسناد بدون نتیجه", style="Heading 2")
        rows = []
        for d in without:
            explanation = doc_explanation(d.doc)
            if d.rows is None and d.doc.status in ("succeeded", "reused"):
                explanation = " ".join(filter(None, [explanation, "نتیجه این سند دیگر در دسترس نیست."]))
            rows.append([_file(d.doc.file_name), doc_status_label(d.doc), _text(explanation)])
        _table(document, ["فایل", "وضعیت", "توضیح"], rows, [4.6, 3.0, 9.4])
    _paragraph(document, NOT_IN_EXCEL_NOTE, muted=True)


def _topic(document, heading: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    document.add_paragraph(heading, style="Heading 3")
    if rows:
        _table(document, [DOC_COLUMN[0], *headers], rows, [DOC_COLUMN[1], *widths])
    else:
        _paragraph(document, NONE_FOUND, muted=True)


def _company_section(document, company: _Company, several_models: bool) -> None:
    """All documents of one company, topic by topic, rows in chronological order (D-17.6)."""
    document.add_paragraph(company.name, style="Heading 2")
    label = company.labels

    headers = ["فایل", "دوره مالی", "حسابرس", "نوع اظهارنظر", "ریسک کلی"] + (["مدل"] if several_models else [])
    widths = [3.4, 3.4, 3.0, 2.4, 2.0] if not several_models else [2.8, 2.8, 2.5, 2.2, 1.8, 2.0]
    _topic(document, "اسناد", headers, [
        [label[id(d)], _file(d.doc.file_name), _text(_summary(d).get("fiscal_period")),
         _text(_summary(d).get("auditor_name")), _text(_summary(d).get("opinion_type")),
         _text(_summary(d).get("overall_risk"))] + ([_text(f"{d.provider} › {d.model}")] if several_models else [])
        for d in company.docs
    ], widths)

    for heading, column in (("نکات کلیدی و نتیجه‌گیری", "key_points"), ("جزییات سطح ریسک", "risk_details")):
        _topic(document, heading, ["متن"], [[label[id(d)], _text(_summary(d).get(column))]
                                            for d in company.docs if _text(_summary(d).get(column))], [15.1])

    _topic(document, "بند اظهارنظر", ["نوع", "خلاصه دلایل"], [
        [label[id(d)], _text(r.get("opinion_kind")), _text(r.get("reasons_summary"))]
        for d in company.docs for r in d.rows.get("opinion", [])
        if _text(r.get("opinion_kind")) or _text(r.get("reasons_summary"))
    ], [2.8, 12.3])

    for sheet, heading, columns in ITEM_TABLES:
        _topic(document, heading, [name for name, _, _ in columns], [
            [label[id(d)], *(_joined(r, cols) for _, cols, _ in columns)]
            for d in company.docs for r in d.rows.get(sheet, []) if keep_row(sheet, r)
        ], [w for _, _, w in columns])


def build_report(state: RequestState, merge_map: dict[str, str] | None = None) -> bytes:
    """`merge_map`: the downloading user's company merges (member key -> main key), as on the dashboard."""
    docs = collect_docs(state)
    companies = _companies([d for d in docs if d.rows], merge_map or {})
    without = [d for d in docs if not d.rows]
    document = _new_document()
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run(f"{TITLE} · {state.id}")
    _font(footer.runs[0]._r.get_or_add_rPr(), size=SIZES["footer"], color=MUTED)
    _cover(document, state, companies, without)
    several_models = len({(d.provider, d.model) for c in companies for d in c.docs}) > 1
    for company in companies:
        document.add_page_break()
        _company_section(document, company, several_models)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
