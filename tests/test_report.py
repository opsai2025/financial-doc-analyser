"""Word report of a request (D-16.1, round 2: D-17.3–D-17.8)."""

import io

from docx import Document
from test_jobs import env, pdf  # noqa: F401  (shared engine fixture)

from core.models import Role
from services import report, results
from services import providers as pv
from services.dashboard import company_key

BANK, OTHER = "بانک نمونه (سهامی عام)", "بانک دیگر"
LTR = "\u202a{}\u202c".format  # file names keep their own direction ("1398.pdf", not "pdf.1398")


def open_report(state, merge_map=None):
    return Document(io.BytesIO(report.build_report(state, merge_map)))


def header(table) -> list[str]:
    return [c.text for c in table.rows[0].cells]


def column(table, index: int) -> list[str]:
    return [r.cells[index].text for r in table.rows[1:]]


def headings(doc, level: int) -> list[str]:
    return [p.text for p in doc.paragraphs if p.style.name == f"Heading {level}"]


def test_cover_and_documents_without_result(env):
    ali = env.user("ali", Role.ANALYST)
    state = env.run(ali, [pdf("1398.pdf"), pdf("bad.pdf"), pdf("1399.pdf")])
    doc = open_report(state)

    assert doc.paragraphs[0].text == report.TITLE
    assert headings(doc, 2) == ["خلاصه اسناد", "اسناد بدون نتیجه", BANK]  # one section per company
    overview = next(t for t in doc.tables if header(t)[0] == "ردیف")
    assert column(overview, 5) == [LTR("1398.pdf"), LTR("1399.pdf")] and column(overview, 2) == ["۱۳۹۸", "۱۳۹۹"]
    failed = next(t for t in doc.tables if header(t) == ["فایل", "وضعیت", "توضیح"])
    assert column(failed, 0) == [LTR("bad.pdf")] and column(failed, 1) == ["ناموفق"]

    # B Nazanin, sizes + 2 pt (D-17.4); right-to-left paragraphs and tables
    normal = doc.styles["Normal"].element.xml
    assert 'w:cs="B Nazanin"' in normal and 'w:szCs w:val="25"' in normal and "w:bidi" in normal
    assert "w:bidiVisual" in doc.element.body.xml and "w:tblHeader" in doc.element.body.xml
    font_table = next(r.target_part.blob for r in doc.part.rels.values() if r.reltype.endswith("/fontTable"))
    assert b'<w:font w:name="B Nazanin"><w:charset w:val="B2"/>' in font_table  # D-18.1


def test_items_by_company_in_chronological_order(env):
    guest = env.user("gst", Role.GUEST, guest_doc_cap=6)
    pv.save_provider(None, pv.ProviderForm("آزمایشی", "mock", "mock://?delay=0"), env.provider.id)
    names = ["1400.pdf", "1399.pdf", "other 1397.pdf", "1398.pdf", "1399-b.pdf"]
    state = env.run(guest, [pdf(n) for n in names])
    data = {d.file_name: d.data[results.ROOT] for d in state.docs}
    data["other 1397.pdf"][results.PART1]["نام_شرکت"] = OTHER
    # 3.3: basis items carry no «موضوعیت_دارد» of their own; only the section does
    data["1398.pdf"][results.PART2]["بند_مبانی_اظهارنظر"] = {"موضوعیت_دارد": True, "موارد_مطرح_شده": [
        {"شماره_مورد": 1, "عنوان": "مبنای الف", "شرح": "شرح الف", "نوع_دلیل": "سایر"},
        {"شماره_مورد": 2, "عنوان": "مبنای ب", "شرح": "شرح ب", "نوع_دلیل": "سایر"}]}
    data["1400.pdf"][results.PART2]["بند_مبانی_اظهارنظر"] = {"موضوعیت_دارد": False}
    # 3.4: only checklist topics found in the report
    for root in data.values():
        for index, item in enumerate(root[results.PART3]):
            item["در_گزارش_آمده"] = index % 2 == 0
    doc = open_report(state)

    assert headings(doc, 2) == ["خلاصه اسناد", OTHER, BANK]
    documents = [t for t in doc.tables if header(t)[:2] == ["سند", "فایل"]]
    assert column(documents[0], 0) == ["۱۳۹۷"]
    # chronological; two documents of 1399 numbered, ordered by file name (D-17.6, D-17.8)
    assert column(documents[1], 0) == ["۱۳۹۸", "۱۳۹۹ (۱)", "۱۳۹۹ (۲)", "۱۴۰۰"]
    assert column(documents[1], 1) == [LTR("1398.pdf"), LTR("1399-b.pdf"), LTR("1399.pdf"), LTR("1400.pdf")]

    basis = [t for t in doc.tables if header(t)[:3] == ["سند", "شماره", "عنوان"]]
    titles = [(r.cells[0].text, r.cells[2].text) for t in basis for r in t.rows[1:]]
    assert ("۱۳۹۸", "مبنای الف") in titles and ("۱۳۹۸", "مبنای ب") in titles
    assert not any(label == "۱۴۰۰" for label, _ in titles)  # section marked not applicable

    for sheet, _, columns in report.ITEM_TABLES:  # every table: exactly the rows the report keeps, by document
        labels = ["سند", *(name for name, _, _ in columns)]
        shown = [r.cells[0].text for t in doc.tables if header(t) == labels for r in t.rows[1:]]
        expected = sum(1 for d in state.docs for row in results.prepare(d.data).rows[sheet]
                       if report.keep_row(sheet, row))
        assert len(shown) == expected, sheet
    checklist = next(t for t in doc.tables if header(t)[:2] == ["سند", "موضوع"])
    assert column(checklist, 0) == sorted(column(checklist, 0), key=lambda label: label[:4])  # chronological

    merged = open_report(state, {company_key(OTHER): company_key(BANK)})  # the dashboard's merges (D-17.5)
    assert headings(merged, 2) == ["خلاصه اسناد", BANK]
    assert column(next(t for t in merged.tables if header(t)[:2] == ["سند", "فایل"]), 0)[0] == "۱۳۹۷"


def test_guest_report_from_memory_without_control_characters(env):
    guest = env.user("gst", Role.GUEST, guest_doc_cap=3)
    pv.save_provider(None, pv.ProviderForm("آزمایشی", "mock", "mock://?delay=0"), env.provider.id)
    state = env.run(guest, [pdf("1398.pdf")])
    state.docs[0].data[results.ROOT][results.PART1]["نکات_کلیدی_و_نتیجه_گیری"] = "نکته\x07 مهم\nسطر دوم"
    doc = open_report(state)
    key_points = [r.cells[1].text for t in doc.tables if header(t) == ["سند", "متن"] for r in t.rows[1:]]
    assert "نکته مهم\nسطر دوم" in key_points
    assert report.report_name(state) == f"{state.id}.docx"
