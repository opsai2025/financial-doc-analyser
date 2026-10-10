"""Integrated Excel workbook (D-8, D-12.9–D-12.13)."""

import io

from openpyxl import load_workbook
from test_jobs import env, pdf  # noqa: F401  (shared engine fixture)

from core.models import Role
from core.sheets import SHEETS
from services import excel, results
from services import providers as pv


def open_book(state):
    return load_workbook(io.BytesIO(excel.build_workbook(state)))


def headers(ws):
    return [c.value for c in ws[1]]


def test_sheets_columns_and_status(env):
    ali = env.user("ali", Role.ANALYST)
    state = env.run(ali, [pdf("1398.pdf"), pdf("bad.pdf"), pdf("1399.pdf")])
    book = open_book(state)
    assert book.sheetnames == [excel.STATUS_SHEET, *(s.title for s in SHEETS)]  # v1 sheet names (D-8.2)
    for spec in SHEETS:
        ws = book[spec.title]
        assert headers(ws) == [*excel.ID_COLUMNS, *(p for _, p in spec.columns)]  # D-12.10, no extra column
        assert ws.sheet_view.rightToLeft and ws.freeze_panes == "A2" and ws.auto_filter.ref
    summary = book["بخش1_خلاصه"]
    assert summary.max_row == 3  # two successful documents, the failed one is only in the status sheet
    assert [c.value for c in summary[2]][:4] == [state.id, state.docs[0].document_id[:12], "بانک نمونه (سهامی عام)",
                                                  "1398.pdf"]
    assert summary.cell(2, headers(summary).index("year") + 1).value == 1398

    status = book[excel.STATUS_SHEET]
    assert headers(status) == list(excel.STATUS_COLUMNS)
    rows = [dict(zip(headers(status), (c.value for c in row), strict=True)) for row in status.iter_rows(min_row=2)]
    assert [r["نام_فایل"] for r in rows] == ["1398.pdf", "bad.pdf", "1399.pdf"]
    assert [r["وضعیت"] for r in rows] == ["موفق", "ناموفق", "موفق"]
    assert rows[0]["SHA-256"] == state.docs[0].document_id and rows[0]["مدل"] == "mock-1"
    assert rows[0]["ذخیره_در_پایگاه"] == "بله" and rows[0]["PDF_نگه‌داری_شد"] == "بله"
    assert rows[1]["توضیح"] and rows[1]["زمان_تحلیل"] in (None, "")


def test_guest_workbook_from_memory(env):
    guest = env.user("gst", Role.GUEST, guest_doc_cap=3)
    state = env.run(guest, [pdf("1398.pdf")])
    book = open_book(state)
    assert book["بخش3_چک_لیست"].max_row > 1
    status = book[excel.STATUS_SHEET]
    row = {h: c.value for h, c in zip(headers(status), status[2], strict=True)}
    assert row["ذخیره_در_پایگاه"] == "خیر" and row["PDF_نگه‌داری_شد"] == "خیر"


def test_long_text_unexpected_fields_and_control_characters(env, monkeypatch):
    """D-12.12/D-12.13: cut only in Excel with a marker and a note; unexpected fields in «سایر_فیلدها»."""
    guest = env.user("gst", Role.GUEST, guest_doc_cap=3)
    pv.save_provider(None, pv.ProviderForm("آزمایشی", "mock", "mock://?delay=0"), env.provider.id)
    state = env.run(guest, [pdf("1398.pdf")])
    data = state.docs[0].data
    report = data[results.ROOT]
    report[results.PART1]["نکات_کلیدی_و_نتیجه_گیری"] = "الف" * 40_000 + "\x07"
    report[results.PART3].append({"گروه": {"موضوع": "تو در تو"}})
    book = open_book(state)

    summary = book["بخش1_خلاصه"]
    cell = summary.cell(2, headers(summary).index("نکات_کلیدی_و_نتیجه_گیری") + 1).value
    assert len(cell) == excel.CELL_LIMIT and cell.endswith(excel.CUT_MARKER) and "\x07" not in cell
    assert len(data[results.ROOT][results.PART1]["نکات_کلیدی_و_نتیجه_گیری"]) == 120_001  # source untouched

    checklist = book["بخش3_چک_لیست"]
    assert headers(checklist)[-1] == excel.EXTRA_COLUMN
    assert "تو در تو" in checklist.cell(checklist.max_row, checklist.max_column).value
    assert excel.EXTRA_COLUMN not in headers(book["بند_اظهارنظر"])

    status = book[excel.STATUS_SHEET]
    note = status.cell(2, list(excel.STATUS_COLUMNS).index("توضیح") + 1).value
    assert "نکات_کلیدی_و_نتیجه_گیری" in note
